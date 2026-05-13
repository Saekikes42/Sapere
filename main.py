import os
import hashlib
import requests
import json
import re
import base64
import pandas as pd
import numpy as np
import streamlit as st
import io
import PyPDF2
from PIL import Image
from docx import Document as DocxDocument
from pptx import Presentation
from sqlalchemy import create_engine, text, types
from datetime import datetime
from bs4 import BeautifulSoup
import altair as alt
from urllib.parse import urljoin

# --- CONFIG ---
st.set_page_config(
    page_title="Sapere Intelligence",
    page_icon="docx",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Modern Enterprise CSS (Matte Dark)
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    
    .stApp {
        background-color: #0d1117;
        color: #c9d1d9;
        font-family: 'Inter', -apple-system, sans-serif;
    }
    
    /* Modern Headers */
    h1 {
        color: #f9a03f !important;
        font-size: 52px !important;
        font-weight: 700 !important;
        margin-bottom: 30px !important;
        letter-spacing: -0.02em;
    }
    
    h2 {
        color: #f9a03f !important;
        font-size: 28px !important;
        font-weight: 600 !important;
        margin-top: 30px !important;
        margin-bottom: 20px !important;
    }
    
    /* Clean Tabs */
    .stTabs [data-baseweb="tab-list"] {
        gap: 4px;
        background-color: #161b22;
        padding: 6px;
        border-radius: 8px;
        border: 1px solid #30363d;
    }

    .stTabs [data-baseweb="tab"] {
        border-radius: 6px;
        padding: 8px 20px;
        color: #8b949e !important;
        transition: all 0.2s ease;
        border: none !important;
    }

    .stTabs [aria-selected="true"] {
        background-color: #21262d !important;
        color: #f9a03f !important;
        font-weight: 600 !important;
    }

    /* Solid Enterprise Buttons */
    .stButton>button {
        background-color: #21262d;
        color: #f9a03f !important;
        border: 1px solid #30363d !important;
        border-radius: 6px !important;
        font-weight: 600 !important;
        font-size: 18px !important;
        width: 100%;
        height: 3.2em !important;
        transition: all 0.1s ease-in-out;
    }
    
    .stButton>button:hover {
        background-color: #30363d;
        border-color: #f9a03f !important;
    }
    
    .stButton>button:active {
        background-color: #f9a03f;
        color: #0d1117 !important;
    }
    
    /* Professional Metrics */
    [data-testid="stMetricValue"] {
        color: #f9a03f !important;
        font-size: 38px !important;
        font-weight: 700 !important;
    }
    
    [data-testid="stMetricLabel"] {
        color: #8b949e !important;
        font-size: 15px !important;
        text-transform: uppercase;
        font-weight: 500;
    }
    
    /* Card Styling */
    .stAlert {
        border-radius: 8px;
        border: 1px solid #30363d;
        background-color: #161b22;
        color: #c9d1d9;
    }
    
    .stDataFrame {
        border: 1px solid #30363d;
        border-radius: 8px;
    }
    
    /* Expander Layout */
    .streamlit-expanderHeader {
        background-color: #161b22 !important;
        border: 1px solid #30363d !important;
        border-radius: 8px !important;
        color: #f9a03f !important;
        font-size: 18px !important;
    }
    
    .block-container {
        padding-top: 3rem !important;
        max-width: 1200px !important;
    }

    /* Input Fields */
    .stTextInput input {
        background-color: #0d1117 !important;
        border: 1px solid #30363d !important;
        border-radius: 6px !important;
        color: #c9d1d9 !important;
        font-size: 17px !important;
    }
    </style>
    """, unsafe_allow_html=True)

# --- CONFIG & DB (Hardcoded) ---
DB_URL = "postgresql+pg8000://postgres:1234@localhost:5432/postgres"
AI_API_KEY = "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"
engine = create_engine(DB_URL)

DOWNLOAD_DIR = os.path.join(os.getcwd(), "downloaded_docs")
if not os.path.exists(DOWNLOAD_DIR): os.makedirs(DOWNLOAD_DIR)

TARGET_URL = "https://cchgeu.ru/education/programms/planned/rupo-9kl-4/?docs2026"
TEXT_MODEL = "google/gemini-2.0-flash-lite-001"

# --- PROMPTS ---
METHODOLOGY_PROMPT = """Оцени ЭОМ (электронный образовательный материал) по критериям:
АВТОМАТИЧЕСКИЕ: описание заполнено; указаны темы/дидактические единицы; >=6 фрагментов; экраны заполнены; шрифт>=22пт; заимствование<=20%; текст не выходит за слайд; названия и время фрагментов указаны.
НОРМАТИВНО-ПРАВОВЫЕ: соответствие ФГОС; защита детей (436-ФЗ); нет противоречий Конституции РФ; нет экстремизма; актуальные НПА; корректные карты/границы РФ; соответствие ГК РФ ст.1274,1259; гражданская идентичность.
НОРМАТИВНО-ТЕХНИЧЕСКИЕ: соответствие ТУ/регламентам; актуальные стандарты и СанПиН.
МЕТОДОЛОГИЧЕСКИЕ: последовательность и завершённость; логичность; практико-ориентированность; критерии оценивания; практические задания; современные тенденции; потребности рынка труда; новые технологии; междисциплинарность; экология; критическое мышление; активное обучение; деятельностный подход.
МЕТОДИЧЕСКИЕ: соответствие темам каркаса; корректная терминология; научная достоверность; однозначность формулировок; указание источников; корректные спецсимволы; нормы речи; соответствие названия содержанию; возрастное соответствие; доп.источники; полнота по ФРП; познавательная ценность; уникальность фрагментов.
ТЕХНИЧЕСКИЕ: нет водяных знаков/рекламы; чёткое аудио/видео; единообразие названий; единый стиль; пастельный фон; не перегружены; нет сторонних ссылок; контент по зонам (доска/планшет учителя/планшет ученика).
Верни JSON: {"compliance_rate":0-1,"structure":0-5,"clarity":0-5,"terminology":0-5,"completeness":0-5,"practice_orientation":0-5,"legal_compliance":0-5,"technical_quality":0-5,"subject":"","topic":"","annotation":"","material_type":"","issues":[],"recommendations":[],"usage_allowed": "yes/no"}"""

# --- CORE LOGIC ---
@st.experimental_singleton(show_spinner=False)
def get_embedding_model():
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer('intfloat/multilingual-e5-small')

def calculate_thematic_links(df):
    if df.empty or len(df) < 2: return df
    try:
        model = get_embedding_model()
        texts = [f"passage: {r.get('topic','')} {r.get('annotation','')}" for _, r in df.iterrows()]
        embs = model.encode(texts, convert_to_numpy=True)
        n = len(df)
        sims = (embs / (np.linalg.norm(embs, axis=1, keepdims=True) + 1e-12)) @ (embs / (np.linalg.norm(embs, axis=1, keepdims=True) + 1e-12)).T
        path = [0]; visited = {0}
        for _ in range(n-1):
            last = path[-1]
            candidates = [i for i in range(n) if i not in visited]
            next_idx = candidates[np.argmax(sims[last, candidates])]
            path.append(next_idx); visited.add(next_idx)
        order = np.argsort(path)
        df['embedding_order'] = order
        df['has_previous'] = df['embedding_order'] > 0
        df['has_next'] = df['embedding_order'] < (n-1)
    except: pass
    return df

def get_document_links(url):
    try:
        r = requests.get(url, timeout=30)
        soup = BeautifulSoup(r.text, 'html.parser')
        links = []
        for a in soup.find_all('a', href=True):
            href = a['href']
            if any(ext in href.lower() for ext in ['.pdf', '.docx', '.pptx', '.doc', '.png', '.jpg', '.csv', '.xlsx']):
                if not href.lower().endswith('.sig'):
                    links.append({'Название': os.path.basename(href), 'url': urljoin(url, href)})
        return links
    except: return []

def download_file(url, dest):
    fname = os.path.basename(url)
    fpath = os.path.join(dest, fname)
    if os.path.exists(fpath): return fpath, False
    try:
        r = requests.get(url, stream=True, timeout=60)
        with open(fpath, 'wb') as f:
            for c in r.iter_content(8192): f.write(c)
        return fpath, True
    except: return None, False

def extract_content(fpath):
    ext = os.path.splitext(fpath)[1].lower()
    text, imgs = "", 0
    try:
        if ext == '.docx':
            d = DocxDocument(fpath); text = "\n".join([p.text for p in d.paragraphs]); imgs = len(d.inline_shapes)
        elif ext == '.pdf':
            with open(fpath, 'rb') as f:
                pdf = PyPDF2.PdfReader(f)
                for p in pdf.pages: text += (p.extract_text() or "") + "\n"
        elif ext == '.pptx':
            p = Presentation(fpath)
            for s in p.slides:
                for sh in s.shapes:
                    if hasattr(sh, "text"): text += sh.text + "\n"
                    if sh.shape_type == 13: imgs += 1
        elif ext in ['.csv', '.xlsx']:
            df = pd.read_csv(fpath) if ext == '.csv' else pd.read_excel(fpath); text = df.to_string()
    except: pass
    return text, imgs

def call_ai(prompt, is_json=False):
    h = {"Authorization": f"Bearer {AI_API_KEY}"}
    d = {"model": TEXT_MODEL, "messages": [{"role": "user", "content": prompt}]}
    try:
        r = requests.post("https://openrouter.ai/api/v1/chat/completions", headers=h, json=d, timeout=60).json()
        c = r['choices'][0]['message']['content'].strip()
        if is_json:
            m = re.search(r'\{.*\}', c, re.DOTALL)
            return json.loads(m.group()) if m else None
        return c
    except: return None

# --- DB INIT ---
def init_db():
    try:
        # 1. Create table if not exists (separate transaction)
        with engine.begin() as conn:
            conn.execute(text('''
                CREATE TABLE IF NOT EXISTS opit_files (
                    id TEXT PRIMARY KEY, 
                    file_name TEXT, 
                    processed_at TIMESTAMP
                )
            '''))
        
        # 2. Get existing columns
        with engine.connect() as conn:
            existing_cols = conn.execute(text(
                "SELECT column_name FROM information_schema.columns WHERE table_name = 'opit_files'"
            )).fetchall()
            existing_cols = [c[0] for c in existing_cols]
        
        # 3. Add missing columns (one by one)
        columns = [
            ("file_path", "TEXT"), ("extension", "TEXT"), ("file_size", "BIGINT"),
            ("text_content", "TEXT"), ("text_length", "INTEGER"), ("images_count", "INTEGER"),
            ("image_analysis", "TEXT"), ("subject", "TEXT"), ("topic", "TEXT"),
            ("annotation", "TEXT"), ("compliance_rate", "REAL"), ("compliance_structure", "REAL"),
            ("compliance_clarity", "REAL"), ("compliance_terminology", "REAL"),
            ("compliance_completeness", "REAL"), ("practice_orientation", "REAL"),
            ("legal_compliance", "REAL"), ("technical_quality", "REAL"),
            ("usage_allowed", "TEXT"), ("issues", "TEXT"), ("recommendations", "TEXT"),
            ("embedding_order", "INTEGER"), ("has_previous", "BOOLEAN"), ("has_next", "BOOLEAN")
        ]
        
        for col_name, col_type in columns:
            if col_name not in existing_cols:
                with engine.begin() as conn:
                    conn.execute(text(f"ALTER TABLE opit_files ADD COLUMN {col_name} {col_type}"))
    except Exception as e:
        st.error(f"Database initialization failed: {e}")
init_db()

# --- REPOSITORY CONFIG ---
if 'repo_path' not in st.session_state:
    st.session_state['repo_path'] = os.path.join(os.getcwd(), "downloaded_docs")

# --- UI ---
st.title("Sapere Intelligence")
tabs = st.tabs(["Загрузка материалов", "Аналитический цикл", "Репозиторий", "Набор данных"])

with tabs[0]:
    st.header("Управление хранилищем")
    
    # Custom Path Selection
    custom_path = st.text_input("Локальный путь к репозиторию", value=st.session_state['repo_path'])
    if custom_path != st.session_state['repo_path']:
        st.session_state['repo_path'] = custom_path
        if not os.path.exists(custom_path):
            try:
                os.makedirs(custom_path)
                st.success(f"Создана новая директория: {custom_path}")
            except Exception as e:
                st.error(f"Ошибка доступа к пути: {e}")
    
    repo_dir = st.session_state['repo_path']

    # --- FOLDER STATISTICS ---
    if os.path.exists(repo_dir):
        files_data = []
        for f in os.listdir(repo_dir):
            fp = os.path.join(repo_dir, f)
            if os.path.isfile(fp):
                stat = os.stat(fp)
                ext = os.path.splitext(f)[1].lower() or "no_ext"
                files_data.append({
                    "Имя": f,
                    "Расширение": ext,
                    "Размер (МБ)": round(stat.st_size / (1024 * 1024), 3),
                    "Изменен": datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M')
                })
        
        if files_data:
            df_files = pd.DataFrame(files_data)
            st.subheader("Технический аудит репозитория")
            
            # Metrics
            m1, m2, m3 = st.columns(3)
            m1.metric("Всего файлов", len(df_files))
            m2.metric("Общий вес", f"{df_files['Размер (МБ)'].sum():.2f} МБ")
            m3.metric("Типов файлов", df_files['Расширение'].nunique())
            
            # Extension Breakdown
            ext_stats = df_files.groupby("Расширение").agg(
                Кол_во=("Имя", "count"),
                Общий_размер_МБ=("Размер (МБ)", "sum")
            ).reset_index()
            
            col_l, col_r = st.columns([1, 1])
            with col_l:
                st.write("**Детализация по типам**")
                st.dataframe(ext_stats, height=300)
            with col_r:
                # Spacer to push chart below the header level of col_l
                st.write("")
                st.write("")
                st.write("")
                # High-End Dashboard Style Chart
                base = alt.Chart(ext_stats).encode(
                    y=alt.Y('Расширение:N', sort='-x', title=None, axis=alt.Axis(labelFontSize=14, labelColor='#8b949e')),
                    x=alt.X('Кол_во:Q', title=None, axis=None)
                )

                bars = base.mark_bar(
                    color='#ff8c00',
                    cornerRadiusEnd=6,
                    height=25
                ).encode(
                    tooltip=['Расширение', 'Кол_во']
                )

                labels = base.mark_text(
                    align='left',
                    baseline='middle',
                    dx=10,
                    color='#f1f5f9',
                    fontSize=15,
                    fontWeight=600
                ).encode(
                    text='Кол_во:Q'
                )

                chart = (bars + labels).properties(
                    height=300,
                    width='container'
                ).configure_view(
                    strokeWidth=0
                ).configure_axis(
                    grid=False
                )
                
                st.altair_chart(chart, use_container_width=True)
            
            with st.expander("Посмотреть полный список файлов в папке"):
                st.table(df_files)
        else:
            st.info("Папка пуста")

    # Manual Upload Section
    with st.expander("Загрузить файлы вручную с компьютера"):
        uploaded_files = st.file_uploader("Выберите файлы для добавления в репозиторий", accept_multiple_files=True)
        if uploaded_files:
            if st.button("Сохранить загруженные файлы"):
                saved_count = 0
                for f in uploaded_files:
                    fpath = os.path.join(repo_dir, f.name)
                    with open(fpath, "wb") as buffer:
                        buffer.write(f.getbuffer())
                    saved_count += 1
                st.success(f"Сохранено в репозиторий: {saved_count} файл(ов)")

    st.markdown("---")
    st.header("Сбор учебного контента с сайта")
    target_url = st.text_input("Введите URL для сканирования", value="https://cchgeu.ru/education/programms/planned/rupo-9kl-4/?docs2026")
    
    if st.button("Начать сканирование сайта"):
        if target_url:
            st.session_state['web_links'] = get_document_links(target_url)
        else:
            st.error("Пожалуйста, введите корректный URL")
    
    if 'web_links' in st.session_state:
        links = st.session_state['web_links']
        st.dataframe(pd.DataFrame(links), height=300)
        if st.button(f"Скачать в репозиторий ({len(links)} ед.)"):
            prog = st.progress(0); status = st.empty(); count = 0
            for i, l in enumerate(links):
                status.text(f"Процесс: {l['Название']}")
                _, new = download_file(l['url'], repo_dir)
                if new: count += 1
                prog.progress((i+1)/len(links))
            st.success(f"Загрузка завершена. Новых материалов: {count}")

with tabs[1]:
    st.header("Аналитический цикл")
    repo_dir = st.session_state['repo_path']
    all_files = [f for f in os.listdir(repo_dir)] if os.path.exists(repo_dir) else []
    
    if not all_files: 
        st.warning("Хранилище пусто. Загрузите файлы на первой вкладке.")
    else:
        # Section 1: AI Moderation with Interactivity
        st.subheader("1. Интеллектуальная модерация (OpenRouter)")
        st.info("Анализ контента, извлечение тем и формирование справок.")
        
        # Identify available extensions
        available_exts = sorted(list(set([os.path.splitext(f)[1].lower() for f in all_files if os.path.splitext(f)[1]])))
        
        col_f1, col_f2 = st.columns([2, 1])
        with col_f1:
            selected_exts = st.multiselect("Выберите типы файлов для анализа", available_exts, default=available_exts)
        with col_f2:
            limit_files = st.number_input("Лимит файлов (0 - без лимита)", min_value=0, value=0)

        # Filter the file list
        files_to_process = [f for f in all_files if os.path.splitext(f)[1].lower() in selected_exts]
        if limit_files > 0:
            files_to_process = files_to_process[:limit_files]

        st.write(f"Объектов к обработке: **{len(files_to_process)}** (отфильтровано из {len(all_files)})")
        
        if st.button("Запустить AI-модерацию выбранных файлов"):
            prog = st.progress(0); status = st.empty()
            for i, fname in enumerate(files_to_process):
                fpath = os.path.join(repo_dir, fname); status.text(f"Анализ ({i+1}/{len(files_to_process)}): {fname}")
                txt, imgs = extract_content(fpath)
                
                # Full AI Analysis with detailed metrics
                prompt = f"{METHODOLOGY_PROMPT}\n\nТЕКСТ ДЛЯ ОЦЕНКИ:\n{txt[:3000]}"
                res = call_ai(prompt, is_json=True)
                
                row = {
                    "id": hashlib.md5(fname.encode()).hexdigest(), 
                    "file_name": fname, 
                    "extension": os.path.splitext(fname)[1], 
                    "text_content": txt,
                    "file_path": fpath,
                    "processed_at": datetime.now()
                }
                
                if res:
                    row.update({
                        "subject": res.get('subject', ''),
                        "topic": res.get('topic', ''),
                        "annotation": res.get('annotation', ''),
                        "compliance_rate": float(res.get('compliance_rate', 0)),
                        "compliance_structure": float(res.get('structure', 0)),
                        "compliance_clarity": float(res.get('clarity', 0)),
                        "compliance_terminology": float(res.get('terminology', 0)),
                        "compliance_completeness": float(res.get('completeness', 0)),
                        "practice_orientation": float(res.get('practice_orientation', 0)),
                        "legal_compliance": float(res.get('legal_compliance', 0)),
                        "technical_quality": float(res.get('technical_quality', 0)),
                        "usage_allowed": res.get('usage_allowed', 'no'),
                        "issues": json.dumps(res.get('issues', []), ensure_ascii=False),
                        "recommendations": json.dumps(res.get('recommendations', []), ensure_ascii=False)
                    })
                
                with engine.begin() as conn:
                    conn.execute(text("DELETE FROM opit_files WHERE id = :id"), {"id": row['id']})
                    pd.DataFrame([row]).to_sql("opit_files", conn, if_exists="append", index=False)
                prog.progress((i+1)/len(files_to_process))
            st.success("Модерация выбранных файлов завершена.")

        st.markdown("---")
        
        # Section 2: Semantic Linkage
        st.subheader("2. Семантическое связывание (Эмбеддинги)")
        st.info("Расчет тематической близости и выстраивание логической последовательности.")
        
        if st.button("Обновить логические связи в базе"):
            try:
                df_final = pd.read_sql("SELECT * FROM opit_files", engine)
                if df_final.empty:
                    st.error("База данных пуста. Сначала проведите модерацию.")
                else:
                    with st.spinner("Рассчитываем эмбеддинги..."):
                        df_linked = calculate_thematic_links(df_final)
                        for _, r in df_linked.iterrows():
                            with engine.begin() as conn:
                                conn.execute(text("UPDATE opit_files SET embedding_order=:o, has_previous=:p, has_next=:n WHERE id=:id"),
                                             {"o":int(r['embedding_order']), "p":bool(r['has_previous']), "n":bool(r['has_next']), "id":r['id']})
                        st.success("Семантическая структура репозитория обновлена.")
            except Exception as e:
                st.error(f"Ошибка при связывании: {e}")

with tabs[2]:
    st.header("Репозиторий объектов")
    search = st.text_input("Поиск по базе знаний (названия, контент, аннотации)")
    try:
        df_repo = pd.read_sql("SELECT * FROM opit_files ORDER BY embedding_order", engine)
        if not df_repo.empty:
            if search: 
                mask = (df_repo['text_content'].str.contains(search, case=False, na=False) | 
                        df_repo['annotation'].str.contains(search, case=False, na=False) |
                        df_repo['file_name'].str.contains(search, case=False, na=False))
                df_repo = df_repo[mask]
            
            for _, r in df_repo.iterrows():
                with st.expander(f"{r['file_name']} | Предмет: {r['subject'] or 'Не определен'} | Рейтинг: {int(r['compliance_rate']*100)}%"):
                    st.markdown(f"### Тема: {r['topic']}")
                    st.write(f"**Аннотация:** {r['annotation']}")
                    
                    c1, c2 = st.columns(2)
                    with c1:
                        if not r['has_previous'] and r['has_next']:
                            st.warning("Статус: Вводный материал (Начало курса)")
                        elif r['has_previous'] and r['has_next']:
                            st.info("Статус: Промежуточный материал")
                        elif r['has_previous'] and not r['has_next']:
                            st.success("Статус: Завершающий материал (Конец курса)")
                        else:
                            st.write("Логические связи не определены")
                    
                    with c2:
                        st.write(f"**Допуск к использованию:** {'ДА' if r['usage_allowed']=='yes' else 'НЕТ'}")
                    
                    # Formatted Lists
                    col_iss, col_rec = st.columns(2)
                    with col_iss:
                        st.markdown("**Критические замечания:**")
                        try:
                            iss = json.loads(r['issues'])
                            if isinstance(iss, list) and iss:
                                for item in iss: st.write(f"• {item}")
                            else: st.write("Замечаний нет")
                        except: st.write(r['issues'])
                    
                    with col_rec:
                        st.markdown("**Методические рекомендации:**")
                        try:
                            recs = json.loads(r['recommendations'])
                            if isinstance(recs, list) and recs:
                                for item in recs: st.write(f"• {item}")
                            else: st.write("Рекомендаций нет")
                        except: st.write(r['recommendations'])
                        
                    if st.checkbox("Показать извлеченный текст документа", key=f"txt_{r['id']}"):
                        st.text_area("Raw Text", value=r['text_content'][:5000], height=300)
        else: 
            st.info("База знаний пуста. Проведите анализ на предыдущей вкладке.")
    except Exception as e:
        st.error(f"Ошибка загрузки репозитория: {e}")

with tabs[3]:
    st.header("Статистика и Экспорт")
    try:
        df_stat = pd.read_sql("SELECT * FROM opit_files", engine)
        if not df_stat.empty:
            m1, m2, m3 = st.columns(3)
            m1.metric("Всего", len(df_stat)); m2.metric("Ср. Качество", round(df_stat['compliance_rate'].mean(), 2)); m3.metric("Допущено", len(df_stat[df_stat['usage_allowed'] == 'yes']))
            st.subheader("Корреляционная матрица")
            st.write(df_stat.select_dtypes(include=[np.number]).corr())
            csv = df_stat.to_csv(index=False).encode('utf-8-sig')
            st.download_button("Выгрузить датасет (CSV)", csv, "dataset.csv", "text/csv")
        else: st.info("Нет данных.")
    except: pass
